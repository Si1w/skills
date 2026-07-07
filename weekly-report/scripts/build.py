"""build.py: assemble weekly-report.pptx from the template and a JSON manifest.

inspect: print every template slide with its shape ids and current text.
         Normally unnecessary: scripts/schema.py catalogs each template by id.
build:   read a manifest (validated against schema.Manifest) that picks
         templates by id (any order, repeats allowed), replaces text and images
         by shape id, and PREPENDS the new slides to the existing output deck.
         The deck accumulates week over week; pass --fresh to start over.

Manifest format (see schema.py for template ids and their fields):
{
  "slides": [
    {"template": "title",
     "text":   {"86": "Project: Foo", "87": "Researcher: Bar | Date: 2026-07-07"}},
    {"template": "results",
     "text":   {"154": "method", "156": "insight"},
     "images": {"151": "/abs/path/fig.png"},
     "delete": ["159"]}
  ]
}
Multi-line text uses "\n".
"""

import argparse
import copy
import io
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches
from PIL import Image

from schema import TEMPLATES, Manifest

SKILL_DIR = Path(__file__).parent.parent
TEMPLATE = SKILL_DIR / "template" / "Weekly Research Update Template.pptx"
R_EMBED = qn("r:embed")
R_LINK = qn("r:link")


def inspect(prs):
    index_to_id = {v["index"]: k for k, v in TEMPLATES.items()}
    for i, slide in enumerate(prs.slides):
        print(f"=== slide {i}: template id {index_to_id.get(i, '?')!r}")
        for sh in slide.shapes:
            kind = str(sh.shape_type).split(" ")[0]
            text = ""
            if sh.has_text_frame:
                text = " / ".join(p.text for p in sh.text_frame.paragraphs if p.text)
            print(f"  id={sh.shape_id:<4} {kind:<10} {text[:120]}")


def set_text(shape, text):
    """Replace a shape's text, keeping the first run's formatting per paragraph."""
    tf = shape.text_frame
    lines = text.split("\n")
    proto = copy.deepcopy(tf.paragraphs[0]._p)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            tf.paragraphs[-1]._p.addnext(copy.deepcopy(proto))
            p = tf.paragraphs[-1]
        runs = p.runs
        if not runs:
            p.text = line
            continue
        runs[0].text = line
        for r in runs[1:]:
            r._r.getparent().remove(r._r)


def shape_map(slide):
    """Map shape id -> shape, recursing into groups so nested shapes are addressable."""
    out = {}

    def walk(shapes):
        for sh in shapes:
            out[str(sh.shape_id)] = sh
            if sh.shape_type == 6:  # GROUP
                walk(sh.shapes)

    walk(slide.shapes)
    return out


def insert_figure(slide, shapes, slot, path):
    """Clear the sample-chart shapes and add the image fitted into the slot region."""
    for sid in slot["clear"]:
        el = shapes[sid]._element
        el.getparent().remove(el)
    left, top, width, height = (Inches(v) for v in slot["region"])
    w, h = Image.open(path).size
    scale = min(width / w, height / h)
    fit_w, fit_h = int(w * scale), int(h * scale)
    slide.shapes.add_picture(path, left + (width - fit_w) // 2, top + (height - fit_h) // 2,
                             fit_w, fit_h)


def set_image(slide, shape, path):
    """Swap the image behind an existing picture shape, keeping its frame."""
    blip = shape._element.find(f".//{qn('a:blip')}")
    if blip is None:
        sys.exit(f"shape id={shape.shape_id} is not a picture")
    _, rid = slide.part.get_or_add_image_part(path)
    blip.set(R_EMBED, rid)


def copy_slide(prs, src):
    """Append a copy of src slide to prs; src may belong to another presentation."""
    layout = next((l for l in prs.slide_layouts if l.name == src.slide_layout.name),
                  prs.slide_layouts[0])
    new = prs.slides.add_slide(layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
            continue
        if rel.is_external:
            rid_map[rid] = new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        elif rel.reltype.endswith("/image"):
            _, rid_map[rid] = new.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
        # other internal rel types don't occur in this template
    for sh in src.shapes:
        el = copy.deepcopy(sh._element)
        for node in el.iter():
            for attr in (R_EMBED, R_LINK):
                old = node.get(attr)
                if old in rid_map:
                    node.set(attr, rid_map[old])
        new.shapes._spTree.append(el)
    return new


def build(manifest, out_path, fresh):
    prs = Presentation(TEMPLATE)
    template_slides = list(prs.slides)

    # claim slides first (duplicating pristine copies for repeats), edit after,
    # so a repeated template slide never inherits an earlier entry's edits
    used = set()
    ordered = []
    for entry in manifest.slides:
        idx = TEMPLATES[entry.template]["index"]
        if idx in used:
            ordered.append(copy_slide(prs, template_slides[idx]))
        else:
            ordered.append(template_slides[idx])
            used.add(idx)

    for entry, slide in zip(manifest.slides, ordered):
        shapes = shape_map(slide)
        for sid, text in entry.text.items():
            set_text(shapes[sid], text)
        for sid, img in entry.images.items():
            if sid == "figure":
                insert_figure(slide, shapes, TEMPLATES[entry.template]["figure_slot"], img)
            else:
                set_image(slide, shapes[sid], img)
        for sid in entry.delete:
            el = shapes[sid]._element
            el.getparent().remove(el)

    # carry over last week's deck below the new slides
    if out_path.exists() and not fresh:
        old = Presentation(out_path)
        for slide in old.slides:
            ordered.append(copy_slide(prs, slide))

    # reorder to `ordered` and drop unused template slides
    sld_id_lst = prs.slides._sldIdLst
    entries = {}
    for sld_id in list(sld_id_lst):
        rid = sld_id.get(qn("r:id"))
        entries[id(prs.part.related_part(rid))] = (sld_id, rid)
        sld_id_lst.remove(sld_id)
    keep = {id(s.part) for s in ordered}
    for part_id, (sld_id, rid) in entries.items():
        if part_id not in keep:
            prs.part.drop_rel(rid)
    for slide in ordered:
        sld_id_lst.append(entries[id(slide.part)][0])

    prs.save(out_path)
    print(f"Wrote {out_path} ({len(ordered)} slides)")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("inspect")
    b = sub.add_parser("build")
    b.add_argument("manifest")
    b.add_argument("-o", "--output", default=str(SKILL_DIR / "weekly-report.pptx"))
    b.add_argument("--fresh", action="store_true", help="discard the existing deck instead of prepending")
    args = ap.parse_args()

    if args.cmd == "inspect":
        inspect(Presentation(TEMPLATE))
    else:
        manifest = Manifest.model_validate(json.loads(Path(args.manifest).read_text()))
        build(manifest, Path(args.output), args.fresh)


if __name__ == "__main__":
    main()
